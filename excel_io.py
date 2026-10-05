"""Read / write the lead tracker stored as a local Excel file (Local_Lead_Tracker.xlsx)."""
import difflib
import re
import shutil
import warnings
from copy import copy
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.cell_range import MultiCellRange

from schema import (AUTO_STAGE, COL_BY, COL_DATE, COL_ID, COL_NAME, COL_NICHE, FIELDS,
                    FIRST_ROW, LAST_ROW, STAGE_RANK)

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

EXCEL_PATH = Path(__file__).parent / "Local_Lead_Tracker.xlsx"
BACKUP_DIR = Path(__file__).parent / "backups"


def _col_idx(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n - 1


# ---------------------------------------------------------------- reading

def read_lists(path=EXCEL_PATH):
    """Dropdown options keyed by Lists column letter, scoring points by Lists row, niche must-haves."""
    ws = load_workbook(path)["Lists"]
    opts = {c: [str(ws[f"{c}{r}"].value) for r in range(2, 8) if ws[f"{c}{r}"].value not in (None, "")]
            for c in "ABCDEFGHIJ"}
    pts = {r: ws[f"M{r}"].value for r in range(2, 16)}
    must = dict(zip(opts["A"], [str(ws[f"B{r}"].value) for r in range(2, 2 + len(opts["A"]))]))
    return opts, pts, must


def read_grid(path=EXCEL_PATH):
    """Leads sheet as a list of rows of strings. The formula columns (R, U, Y, Z) are calculated
    here with the workbook's own rules, because openpyxl does not evaluate formulas."""
    ws = load_workbook(path)["Leads"]
    _opts, pts, must = read_lists(path)
    grid = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=29):
        grid.append(["" if c.value is None
                     else c.value.strftime("%d-%b-%Y") if isinstance(c.value, (date, datetime))
                     else str(c.value) for c in row])
    g = lambda c, r: ws[f"{c}{r}"].value
    for r in range(FIRST_ROW, ws.max_row + 1):
        if not g("C", r):
            continue
        blanks = sum(g(c, r) in (None, "") for c in "LMOPQST")
        status = "Complete" if blanks == 0 else f"{blanks} missing"
        site = g("L", r)
        score = (pts[2] if site == "None" else pts[3] if site in ("Broken", "Outdated / Poor")
                 else pts[4] if site == "Basic" else 0)
        score += pts[5] if (g("H", r) or 0) >= pts[12] else 0
        score += pts[6] if (g("G", r) or 0) >= pts[13] else 0
        score += pts[7] if g("Q", r) == "No" else 0
        score += pts[8] if g("O", r) == "Active" else 0
        score += pts[9] if g("P", r) == "Yes" else 0
        score += pts[10] if g("V", r) == "Yes" else 0
        score += pts[11] if g("X", r) == "Yes" else 0
        prio = ("Audit pending" if status != "Complete"
                else "Hot" if score >= pts[14] else "Warm" if score >= pts[15] else "Cold")
        for col, val in (("R", must.get(g("B", r), "")), ("U", status), ("Y", score), ("Z", prio)):
            grid[r - 1][_col_idx(col)] = str(val)
    return grid


def cell(grid, col, row):
    r, c = row - 1, _col_idx(col)
    return grid[r][c] if r < len(grid) and c < len(grid[r]) else ""


def existing_leads(grid, niche):
    return [(cell(grid, COL_ID, r), cell(grid, COL_NAME, r)) for r in range(FIRST_ROW, len(grid) + 1)
            if cell(grid, COL_NICHE, r) == niche and cell(grid, COL_NAME, r)]


def leads_dataframe(grid):
    headers = [str(h).replace(" *", "") for h in grid[1]]
    rows = [r + [""] * (len(headers) - len(r)) for r in grid[FIRST_ROW - 1:]]
    df = pd.DataFrame(rows, columns=headers).replace("", None)
    score = "Opportunity Score (/10)"
    df[score] = pd.to_numeric(df[score], errors="coerce")
    return df


def dashboard_dataframe(df, target=10):
    score = "Opportunity Score (/10)"
    rows = []
    for niche, g in df.groupby("Niche", sort=False):
        done = g[g["Audit Status"] == "Complete"]
        rows.append({
            "Niche": niche, "Target": target, "Captured": int(g["Business Name"].notna().sum()),
            "Audit Complete": len(done), "Hot": int((g["Priority"] == "Hot").sum()),
            "Warm": int((g["Priority"] == "Warm").sum()), "Cold": int((g["Priority"] == "Cold").sum()),
            "Qualified": int((g["Stage"] == "Qualified").sum()),
            "Disqualified": int((g["Stage"] == "Disqualified").sum()),
            "No Website": int((g["Website Status"] == "None").sum()),
            "Avg Score": round(done[score].mean(), 1) if len(done) else 0})
    out = pd.DataFrame(rows)
    total = {"Niche": "TOTAL", **out.drop(columns=["Niche", "Avg Score"]).sum().to_dict()}
    all_done = df[df["Audit Status"] == "Complete"][score]
    total["Avg Score"] = round(all_done.mean(), 1) if len(all_done) else 0
    out.loc[len(out)] = total
    out["% of Target"] = (out["Captured"] / out["Target"] * 100).round(0).astype(int).astype(str) + "%"
    return out


# ---------------------------------------------------------------- writing

def _convert(ftype, value):
    if value in (None, ""):
        return None
    if ftype == "float":
        return float(value)
    if ftype == "int":
        return int(float(str(value).replace(",", "")))
    return str(value)


# words too common to tell two businesses apart
_GENERIC = {"the", "and", "a", "of", "salon", "studio", "parlour", "parlor", "unisex", "beauty", "ladies",
            "gents", "clinic", "dental", "dentist", "gym", "fitness", "cafe", "restaurant", "turf",
            "football", "family", "kolkata"}


def _distinct(name):
    """The words of a name that tell businesses apart (generic words like 'salon' removed)."""
    words = re.findall(r"[a-z0-9]+", name.casefold())
    return [w for w in words if w not in _GENERIC] or words


def _best_match(name, candidates):
    """Row of the stored name that best matches `name`, or None.
    'Glow Studio' matches 'Glow Studio Unisex Salon' (same distinctive word), but 'Kavya Salon' does not match
    'Jaya Salon': only the distinctive words are compared, so a shared word like 'Salon' cannot make two
    different businesses look alike."""
    key = _distinct(name)
    best, best_score = None, 0.0
    for row, other in candidates.items():
        kn = _distinct(other)
        ratio = difflib.SequenceMatcher(None, " ".join(key), " ".join(kn)).ratio()
        score = max(ratio, 0.9) if set(key) <= set(kn) or set(kn) <= set(key) else ratio
        if score > best_score:
            best, best_score = row, score
    return best if best_score >= 0.85 else None


def plan_rows(grid, niche, leads, selected_id=None, allow_new=True):
    """Decide which sheet row each parsed lead goes to. Actions:
    'update'   an existing lead (matched by the chosen lead or by name)
    'new'      the first free slot of the niche
    'extend'   all slots are full: an extra row is added (SAL-11, SAL-12, ...), reusing an emptied extra row first
    'nomatch'  a Step 2/3 note whose lead is unknown (allow_new=False): nothing is written
    'noname'   no name and no match: nothing is written"""
    last = len(grid)
    rows = [r for r in range(FIRST_ROW, last + 1) if cell(grid, COL_NICHE, r) == niche]
    names = {r: cell(grid, COL_NAME, r).strip() for r in rows}
    spare = [r for r in range(LAST_ROW + 1, last + 1) if not cell(grid, COL_NICHE, r)]  # emptied extra rows
    next_row = last + 1
    ids = [i for i in (cell(grid, COL_ID, r) for r in rows) if re.fullmatch(r".+-\d+", i)]
    prefix = ids[0].rsplit("-", 1)[0] if ids else niche[:3].upper()
    next_num = max((int(i.rsplit("-", 1)[1]) for i in ids), default=0) + 1
    used, plan = set(), []
    for lead in leads:
        row, lead_id = None, None
        if selected_id and len(leads) == 1:
            row = next((r for r in rows if cell(grid, COL_ID, r) == selected_id), None)
        name = str(lead.get("business_name") or "").strip()
        if row is None and name:
            filled = {r: n for r, n in names.items() if n and r not in used}
            row = _best_match(name, filled)
        action = "update"
        if row is None:
            free = [r for r in rows if not names[r] and r not in used]
            if not name:
                action = "noname"
            elif not allow_new:
                action = "nomatch"
            elif free:
                row, action = free[0], "new"
            else:
                action = "extend"
                if spare:
                    row = spare.pop(0)
                else:
                    row, next_row = next_row, next_row + 1
                lead_id = f"{prefix}-{next_num:02d}"
                next_num += 1
        if row is not None:
            used.add(row)
        plan.append({"row": row, "action": action,
                     "lead_id": lead_id or (cell(grid, COL_ID, row) if row else None),
                     "has_name": action != "noname"})
    return plan


def build_updates(grid, niche, leads, steps, researched_by="", selected_id=None):
    """Returns (list of (A1 cell, value), result messages, extra rows to create). Pure function - touches no files."""
    plan = plan_rows(grid, niche, leads, selected_id, allow_new=1 in steps)
    auto_stage = AUTO_STAGE[max(steps)]
    updates, msgs, new_rows = [], [], []
    for lead, p in zip(leads, plan):
        row = p["row"]
        if p["action"] == "nomatch":
            msgs.append(f"Could not tell which lead '{lead.get('business_name')}' is. Nothing written.")
            continue
        if p["action"] == "noname":
            msgs.append("Skipped a lead with no business name.")
            continue
        if p["action"] == "extend":
            new_rows.append(row)
            updates += [(f"{COL_ID}{row}", p["lead_id"]), (f"{COL_NICHE}{row}", niche)]
        for key, (col, _l, _s, ftype) in FIELDS.items():
            if key == "business_name" and p["action"] == "update":
                continue  # keep the stored name; it was only used for matching
            val = _convert(ftype.split(":")[0], lead.get(key))
            if val is not None:
                updates.append((f"{col}{row}", val))
        if p["action"] in ("new", "extend"):
            updates.append((f"{COL_DATE}{row}", date.today()))
            if researched_by:
                updates.append((f"{COL_BY}{row}", researched_by))
        if not lead.get("stage"):  # advance Stage automatically, never backwards
            stage_col = FIELDS["stage"][0]
            if STAGE_RANK.get(auto_stage, 0) > STAGE_RANK.get(cell(grid, stage_col, row), 0):
                updates.append((f"{stage_col}{row}", auto_stage))
        name = cell(grid, COL_NAME, row) or lead.get("business_name")
        verb = "Added" if p["action"] == "new" else "Added (extra row)" if p["action"] == "extend" else "Updated"
        msgs.append(f"{verb} {p['lead_id']} - {name}")
    return updates, msgs, new_rows


BIG_ROW = 1000  # formulas and dropdowns are widened to this row when the sheet first grows


def _widen(match_fn, text):
    return re.sub(match_fn, lambda m: m.group(1) + str(BIG_ROW), text)


def _ensure_extended(wb):
    """Widen the dropdown ranges and the Dashboard formulas from row 52 to BIG_ROW (safe to repeat)."""
    ws = wb["Leads"]
    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange(_widen(r"([A-Z]+3:[A-Z]+)52\b", str(dv.sqref)))
    for row in wb["Dashboard"].iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.value = _widen(r"(\$[A-Z]+\$3:\$[A-Z]+\$)52\b", c.value)


def _init_row(ws, row, template=LAST_ROW):
    """Make a new bottom row look and calculate like the original ones (styles + Score/Priority formulas)."""
    for col in range(1, 30):
        src, dst = ws.cell(template, col), ws.cell(row, col)
        dst._style = copy(src._style)
        if isinstance(src.value, str) and src.value.startswith("="):
            dst.value = Translator(src.value, origin=src.coordinate).translate_formula(dst.coordinate)
    if ws.row_dimensions[template].height:
        ws.row_dimensions[row].height = ws.row_dimensions[template].height


def _backup(path):
    """Timestamped local copy before any change (keeps the newest 20)."""
    BACKUP_DIR.mkdir(exist_ok=True)
    shutil.copy2(path, BACKUP_DIR / f"tracker_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    for old in sorted(BACKUP_DIR.glob("tracker_*.xlsx"))[:-20]:
        old.unlink()


def write_leads(niche, leads, steps, researched_by="", selected_id=None, path=EXCEL_PATH):
    """Apply the parsed leads to the workbook, keeping a timestamped backup. Returns messages."""
    updates, msgs, new_rows = build_updates(read_grid(path), niche, leads, steps, researched_by, selected_id)
    if updates:
        _backup(path)
        wb = load_workbook(path)
        ws = wb["Leads"]
        if new_rows and max(new_rows) > LAST_ROW:
            _ensure_extended(wb)
        for r in sorted(r for r in new_rows if r > ws.max_row):  # brand-new bottom rows (emptied ones already exist)
            _init_row(ws, r)
        for a1, v in updates:
            ws[a1].value = v
        wb.calculation.fullCalcOnLoad = True  # Excel recalculates Score / Priority on open
        wb.save(path)
    return msgs


def _compact(ws, niche):
    """Close the gaps in one niche: leads move up so the filled rows come first (SAL-02 -> SAL-01 ...).
    Lead ID, Niche and the formula columns stay in their rows; only the typed-in data moves."""
    rows = [r for r in range(FIRST_ROW, ws.max_row + 1) if ws[f"{COL_NICHE}{r}"].value == niche]
    cols = range(3, 30)  # C..AC are the input columns
    is_formula = lambda v: isinstance(v, str) and v.startswith("=")
    data = [{c: ws.cell(r, c).value for c in cols if not is_formula(ws.cell(r, c).value)}
            for r in rows if ws[f"{COL_NAME}{r}"].value not in (None, "")]
    for i, r in enumerate(rows):
        for c in cols:
            if not is_formula(ws.cell(r, c).value):
                ws.cell(r, c).value = data[i].get(c) if i < len(data) else None
    for r in reversed(rows):  # an extra row (beyond the original 50) that ended up empty is released for reuse
        if r > LAST_ROW and ws[f"{COL_NAME}{r}"].value in (None, ""):
            ws[f"{COL_ID}{r}"].value = None
            ws[f"{COL_NICHE}{r}"].value = None
        else:
            break


def delete_lead(lead_id, path=EXCEL_PATH):
    """Delete a lead and move the later leads of the same niche up one place. Returns a message."""
    wb = load_workbook(path)
    ws = wb["Leads"]
    row = next((r for r in range(FIRST_ROW, ws.max_row + 1) if ws[f"{COL_ID}{r}"].value == lead_id), None)
    if row is None:
        raise ValueError(f"Lead {lead_id} not found")
    name, niche = ws[f"{COL_NAME}{row}"].value, ws[f"{COL_NICHE}{row}"].value
    later = sum(1 for r in range(row + 1, ws.max_row + 1)
                if ws[f"{COL_NICHE}{r}"].value == niche and ws[f"{COL_NAME}{r}"].value not in (None, ""))
    _backup(path)
    for col in range(3, 30):
        c = ws.cell(row, col)
        if c.value is not None and not str(c.value).startswith("="):
            c.value = None
    _compact(ws, niche)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return f"Deleted {lead_id} - {name}" + (f". The {later} lead(s) after it moved up." if later else "")
